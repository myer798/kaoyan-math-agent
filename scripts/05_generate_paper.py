# -*- coding: utf-8 -*-
"""
第 5 步：出题 Agent。
输入一个知识点（如"多元函数极值"），自动：
  1) 检索题库里最相关的真题作为"范文"
  2) 调 DeepSeek 仿真题风格新编一套模拟卷（选择题5 + 填空题5 + 解答题6）
  3) 为每道新题写详细解析
  4) 同时输出 网页HTML 和 Word .docx

用法： python scripts/05_generate_paper.py "多元函数极值"
"""
import json
import math
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openai import OpenAI  # noqa: E402
import config  # noqa: E402

# ====== 向量检索 ======

def cosine_sim(a: list[float], b: list[float]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(x * x for x in b))
    return dot / (na * nb + 1e-9)


def search_questions(client: OpenAI, topic: str, top_k: int = 12) -> list[dict]:
    """用知识点文本做向量检索，返回最相关的真题。"""
    # 先把知识点编码成向量
    vec_resp = client.embeddings.create(
        model=config.EMBED_MODEL, input=f"知识点:{topic}"
    )
    query_vec = vec_resp.data[0].embedding

    records = json.loads((config.INDEX_DIR / "vectors.json").read_text(encoding="utf-8"))
    scored = []
    for r in records:
        sim = cosine_sim(query_vec, r.get("embedding", []))
        scored.append((sim, r))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [r for _, r in scored[:top_k]]


# ====== 出题提示词 ======

GEN_PROMPT = """你是考研数学（一）命题专家。下面给你一些历年真题作为风格参考：

{examples}

请仿照真题的难度、风格和题型，围绕知识点「{topic}」新编一套模拟卷。
要求：
1. 选择题 5 道（每题 4 个选项 A/B/C/D，每题 5 分）
2. 填空题 5 道（每题 5 分）
3. 解答题 6 道（每题 10-12 分，需写详细过程）
4. 所有数学公式用 LaTeX（行内 $...$，独立 $$...$$）
5. 题目要原创，不是照抄真题，但风格难度要接近
6. 每道题都要写【答案】和【解析】

输出格式为 JSON：
{{
  "title": "试卷标题",
  "questions": [
    {{"qtype":"选择题","qno":"1","content":"...","answer":"A","analysis":"...","topics":["..."]}},
    ...
  ]
}}
只输出 JSON，不要用 ```json 包裹，不要加任何说明。"""


def generate_paper(client: OpenAI, topic: str, examples: list[dict]) -> dict:
    # 把真题范文拼成文本
    ex_text = ""
    for i, q in enumerate(examples, 1):
        ex_text += f"\n【范文{i}】{q['year']}年{q.get('qtype','')}第{q.get('qno','')}题\n"
        ex_text += f"题干: {q.get('content','')[:200]}\n"
        if q.get('analysis'):
            ex_text += f"解析: {q.get('analysis','')[:200]}\n"

    prompt = GEN_PROMPT.format(topic=topic, examples=ex_text)

    for attempt in range(3):
        try:
            resp = client.chat.completions.create(
                model=config.LLM_MODEL,
                messages=[
                    {"role": "system", "content": "你是考研数学命题专家，只输出JSON。"},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.7,
                response_format={"type": "json_object"},
                max_tokens=8000,
            )
            raw = resp.choices[0].message.content.strip()
            obj = json.loads(raw)
            if "questions" in obj:
                return obj
            return {"title": obj.get("title", f"模拟卷：{topic}"), "questions": obj if isinstance(obj, list) else []}
        except Exception as e:  # noqa: BLE001
            print(f"  [重试 {attempt+1}] {str(e)[:100]}")
            if attempt < 2:
                time.sleep(2 ** attempt)
    print("  [警告] 生成失败")
    return {"title": f"模拟卷：{topic}", "questions": []}


# ====== 输出：网页 HTML ======

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<script>
  window.MathJax = {{
    tex: {{ inlineMath: [['$','$'], ['\\\\(','\\\\)']], displayMath: [['$$','$$'],['\\\\[','\\\\]']] }},
    svg: {{ fontCache: 'global' }}
  }};
</script>
<script src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-svg.js"></script>
<style>
  body {{ font-family: "Microsoft YaHei", "SimSun", serif; max-width: 800px; margin: 40px auto; padding: 0 20px; line-height: 1.8; }}
  h1 {{ text-align: center; border-bottom: 2px solid #333; padding-bottom: 10px; }}
  .q {{ margin: 20px 0; padding: 15px 0; border-bottom: 1px dashed #ccc; }}
  .qtype-header {{ font-weight: bold; font-size: 1.1em; margin: 25px 0 10px; }}
  .qno {{ font-weight: bold; }}
  .ans {{ color: #c00; margin-top: 8px; }}
  .ana {{ margin-top: 8px; padding-left: 20px; border-left: 3px solid #4a90d9; color: #444; }}
  .ans::before {{ content: "【答案】"; font-weight: bold; }}
  .ana::before {{ content: "【解析】"; font-weight: bold; color: #4a90d9; }}
</style>
</head>
<body>
<h1>{title}</h1>
{body}
</body>
</html>"""

QTYPE_ORDER = {"选择题": 0, "填空题": 1, "解答题": 2, "证明题": 3}


def to_html(paper: dict) -> str:
    qs = sorted(paper.get("questions", []), key=lambda q: QTYPE_ORDER.get(q.get("qtype", ""), 9))
    body = ""
    current_type = ""
    for q in qs:
        qt = q.get("qtype", "")
        if qt != current_type:
            current_type = qt
            body += f'<div class="qtype-header">{qt}</div>\n'
        body += f'<div class="q">\n<div class="qno">{q.get("qno","")}.</div>\n'
        body += f'<div>{q.get("content","")}</div>\n'
        if q.get("answer"):
            body += f'<div class="ans">{q["answer"]}</div>\n'
        if q.get("analysis"):
            body += f'<div class="ana">{q["analysis"]}</div>\n'
        body += "</div>\n"
    return HTML_TEMPLATE.format(title=paper.get("title", "模拟卷"), body=body)


# ====== 输出：Word .docx ======

def to_docx(paper: dict, out_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    # 标题
    h = doc.add_heading(paper.get("title", "模拟卷"), level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    qs = sorted(paper.get("questions", []), key=lambda q: QTYPE_ORDER.get(q.get("qtype", ""), 9))
    current_type = ""
    for q in qs:
        qt = q.get("qtype", "")
        if qt != current_type:
            current_type = qt
            doc.add_heading(qt, level=1)
        doc.add_paragraph(f'{q.get("qno","")}. {q.get("content","")}')
        if q.get("answer"):
            doc.add_paragraph(f'【答案】{q["answer"]}')
        if q.get("analysis"):
            p = doc.add_paragraph(f'【解析】{q["analysis"]}')
            # 公式不渲染，保持 LaTeX 源码（Word 可后续用插件转换）

    doc.save(str(out_path))


# ====== 主流程 ======

def main() -> None:
    if len(sys.argv) < 2:
        print('用法: python scripts/05_generate_paper.py "知识点"')
        print('示例: python scripts/05_generate_paper.py "多元函数极值"')
        sys.exit(1)
    topic = sys.argv[1]
    print(f"知识点: {topic}")
    print("1) 检索相关真题 ...")

    client_v = OpenAI(api_key=config.VISION_API_KEY, base_url=config.VISION_BASE_URL)
    client_l = OpenAI(api_key=config.LLM_API_KEY, base_url=config.LLM_BASE_URL)

    examples = search_questions(client_v, topic, top_k=12)
    print(f"   找到 {len(examples)} 道相关真题作为范文")
    for e in examples[:5]:
        print(f"   - {e['year']}年 {e.get('qtype','')}第{e.get('qno','')}题  知识点:{e.get('topics',[])}")

    print("2) 调 DeepSeek 生成模拟卷 ...")
    paper = generate_paper(client_l, topic, examples)
    n = len(paper.get("questions", []))
    print(f"   生成 {n} 道题")
    if n == 0:
        print("   生成失败，请重试或换知识点")
        return

    # 存 JSON
    json_path = config.PAPER_DIR / f"paper_{topic[:10]}.json"
    json_path.write_text(json.dumps(paper, ensure_ascii=False, indent=2), encoding="utf-8")

    # 存网页
    html = to_html(paper)
    html_path = config.WEB_DIR / f"paper_{topic[:10]}.html"
    html_path.write_text(html, encoding="utf-8")

    # 存 Word
    docx_path = config.PAPER_DIR / f"paper_{topic[:10]}.docx"
    to_docx(paper, docx_path)

    print("-" * 50)
    print(f"完成！")
    print(f"  JSON: {json_path}")
    print(f"  网页: {html_path}")
    print(f"  Word: {docx_path}")


if __name__ == "__main__":
    main()
