# -*- coding: utf-8 -*-
"""
Web 服务：浏览器输入知识点 → 生成模拟卷 → 页面渲染。
启动： python app.py
然后在浏览器打开 http://127.0.0.1:5000
"""
import json
import math
import sys
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, r"E:\ai_0\pylibs")

from flask import Flask, request, jsonify, send_from_directory  # noqa: E402
from openai import OpenAI  # noqa: E402
import config  # noqa: E402

app = Flask(__name__, static_folder=None)

# ====== 复用 05 脚本里的检索和生成逻辑 ======

def cosine_sim(a, b):
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(x * x for x in b))
    return dot / (na * nb + 1e-9)


def search_questions(client_v, topic, top_k=12):
    vec_resp = client_v.embeddings.create(model=config.EMBED_MODEL, input=f"知识点:{topic}")
    query_vec = vec_resp.data[0].embedding
    records = json.loads((config.INDEX_DIR / "vectors.json").read_text(encoding="utf-8"))
    scored = [(cosine_sim(query_vec, r.get("embedding", [])), r) for r in records]
    scored.sort(key=lambda x: x[0], reverse=True)
    return [r for _, r in scored[:top_k]]


GEN_PROMPT = """你是考研数学（一）命题专家。下面给你一些历年真题作为风格参考：

{examples}

请仿照真题的难度、风格和题型，围绕知识点「{topic}」新编一套模拟卷。
要求：
1. 选择题 5 道（每题 4 个选项 A/B/C/D，每题 5 分）
2. 填空题 5 道（每题 5 分）
3. 解答题 6 道（每题 10-12 分，需写详细过程）
4. 所有数学公式用 LaTeX（行内 $...$，独立 $$...$$）
5. 题目要原创，不是照抄真题，但风格难度要接近
6. 每道题都要写【答案】和【解析】

输出格式为 JSON：
{{
  "title": "试卷标题",
  "questions": [
    {{"qtype":"选择题","qno":"1","content":"...","answer":"A","analysis":"...","topics":["..."]}},
    ...
  ]
}}
只输出 JSON，不要用代码块包裹，不要加任何说明。"""


def generate_paper(client_l, topic, examples):
    ex_text = ""
    for i, q in enumerate(examples, 1):
        ex_text += f"\n【范文{i}】{q['year']}年{q.get('qtype','')}第{q.get('qno','')}题\n"
        ex_text += f"题干: {q.get('content','')[:200]}\n"
        if q.get('analysis'):
            ex_text += f"解析: {q.get('analysis','')[:200]}\n"

    prompt = GEN_PROMPT.format(topic=topic, examples=ex_text)
    for attempt in range(3):
        try:
            resp = client_l.chat.completions.create(
                model=config.LLM_MODEL,
                messages=[
                    {"role": "system", "content": "你是考研数学命题专家，只输出JSON。"},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.7,
                response_format={"type": "json_object"},
                max_tokens=8000,
            )
            raw = resp.choices[0].message.content.strip()
            obj = json.loads(raw)
            if "questions" in obj:
                return obj
            return {"title": obj.get("title", f"模拟卷：{topic}"), "questions": []}
        except Exception as e:  # noqa: BLE001
            if attempt < 2:
                import time; time.sleep(2 ** attempt)
    return {"title": f"模拟卷：{topic}", "questions": []}


# ====== 路由 ======

@app.route("/")
def index():
    return send_from_directory(str(config.ROOT / "web"), "index.html")


@app.route("/api/generate", methods=["POST"])
def api_generate():
    data = request.get_json()
    topic = (data.get("topic") or "").strip()
    if not topic:
        return jsonify({"error": "请输入知识点"}), 400

    client_v = OpenAI(api_key=config.VISION_API_KEY, base_url=config.VISION_BASE_URL)
    client_l = OpenAI(api_key=config.LLM_API_KEY, base_url=config.LLM_BASE_URL)

    # 1) 检索
    examples = search_questions(client_v, topic, top_k=12)
    # 2) 生成
    paper = generate_paper(client_l, topic, examples)

    # 3) 存档
    safe = topic[:10]
    (config.PAPER_DIR / f"paper_{safe}.json").write_text(
        json.dumps(paper, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # 4) 返回 JSON（前端渲染）
    return jsonify({
        "title": paper.get("title", f"模拟卷：{topic}"),
        "examples": [{"year": e["year"], "qtype": e.get("qtype",""), "qno": e.get("qno",""), "topics": e.get("topics",[])} for e in examples],
        "questions": paper.get("questions", []),
    })


@app.route("/api/download/<fmt>/<filename>")
def download(fmt, filename):
    """下载已生成的 Word 文件"""
    if fmt == "docx":
        return send_from_directory(str(config.PAPER_DIR), filename, as_attachment=True)
    return "", 404


@app.route("/api/docx", methods=["POST"])
def api_docx():
    """接收前端 JSON，生成 Word 文件并返回下载"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io

    data = request.get_json()
    title = data.get("title", "模拟卷")
    qs = data.get("questions", [])

    TYPE_ORDER = {"选择题": 0, "填空题": 1, "解答题": 2, "证明题": 3}
    qs = sorted(qs, key=lambda q: TYPE_ORDER.get(q.get("qtype", ""), 9))

    doc = Document()
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cur = ""
    for q in qs:
        if q.get("qtype") != cur:
            cur = q.get("qtype", "")
            doc.add_heading(cur, level=1)
        # 转换 LaTeX 为 Word 可读格式
        content = _latex_to_text(q.get("content", ""))
        doc.add_paragraph(f'{q.get("qno","")}. {content}')
        if q.get("answer"):
            answer = _latex_to_text(q["answer"])
            doc.add_paragraph(f'【答案】{answer}')
        if q.get("analysis"):
            analysis = _latex_to_text(q["analysis"])
            doc.add_paragraph(f'【解析】{analysis}')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_from_directory(
        directory=str(config.PAPER_DIR),
        filename="dummy",  # 占位，实际用 send_file
        as_attachment=True,
    ) if False else __send_buf(buf, title[:20])


def __send_buf(buf, name):
    from flask import send_file
    return send_file(buf, as_attachment=True, download_name=f"模拟卷_{name}.docx")


def _latex_to_text(text):
    """把 LaTeX 公式转换成 Word 能显示的纯文本格式"""
    import re
    # 行内公式 $...$ -> 保留内容
    text = re.sub(r'\$([^\$]+)\$', r'\1', text)
    # 独立公式 $$...$$ -> 保留内容
    text = re.sub(r'\$\$([^\$]+)\$\$', r'\1', text)
    # 常见 LaTeX 命令替换
    replacements = {
        r'\lim\limits': 'lim',
        r'\lim': 'lim',
        r'\sin': 'sin',
        r'\cos': 'cos',
        r'\tan': 'tan',
        r'\dfrac': '',
        r'\frac': '',
        r'\to': '→',
        r'\infty': '∞',
        r'\alpha': 'α',
        r'\beta': 'β',
        r'\gamma': 'γ',
        r'\delta': 'δ',
        r'\epsilon': 'ε',
        r'\theta': 'θ',
        r'\lambda': 'λ',
        r'\mu': 'μ',
        r'\pi': 'π',
        r'\sigma': 'σ',
        r'\phi': 'φ',
        r'\omega': 'ω',
        r'\Delta': 'Δ',
        r'\Sigma': 'Σ',
        r'\Omega': 'Ω',
        r'\in': '∈',
        r'\subset': '⊂',
        r'\cup': '∪',
        r'\cap': '∩',
        r'\leq': '≤',
        r'\geq': '≥',
        r'\neq': '≠',
        r'\approx': '≈',
        r'\equiv': '≡',
        r'\times': '×',
        r'\div': '÷',
        r'\pm': '±',
        r'\cdot': '·',
        r'\circ': '°',
        r'\prime': '′',
        r'\partial': '∂',
        r'\nabla': '∇',
        r'\int': '∫',
        r'\oint': '∮',
        r'\sum': '∑',
        r'\prod': '∏',
        r'\sqrt': '√',
        r'\{': '(',
        r'\}': ')',
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    # 清理多余的花括号
    text = re.sub(r'[{}]', '', text)
    return text


if __name__ == "__main__":
    print("=" * 50)
    print("  数学模拟卷生成系统")
    print("  浏览器打开: http://127.0.0.1:5000")
    print("  按 Ctrl+C 停止")
    print("=" * 50)
    app.run(debug=True, port=5000)
